import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from typing import List, Dict

def get_max_depth(tree_dict, current=0):
    if not tree_dict:
        return current
    return max(get_max_depth(v, current + 1) for v in tree_dict.values() if isinstance(v, dict))

def generate_pubmed_matrix_excel(pubmed_data: dict, final_taxonomy: List[Dict], draft_tree: dict, output_path: str, input_ids: List[str] = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PubMed Taxonomy Matrix"

    fixed_headers = ["S.No", "PMID", "Article Title", "Authors", "Journal", "Abstract"]
    
    fixed_header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark Blue
    fixed_header_font = Font(bold=True, color="FFFFFF")
    
    dynamic_header_font = Font(bold=True, color="000000")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    palette = ["92D050", "00B0F0", "FFC000", "B1A0C7", "FFFF00", "B7DEE8", "D9E1F2", "FCE4D6"]

    max_depth = get_max_depth(draft_tree)
    if max_depth == 0:
        max_depth = 1

    # Write fixed headers
    for col_idx, header in enumerate(fixed_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = fixed_header_font
        cell.fill = fixed_header_fill
        cell.alignment = center_align
        if max_depth > 1:
            ws.merge_cells(start_row=1, start_column=col_idx, end_row=max_depth, end_column=col_idx)
            
    # Set specific widths for PubMed fixed columns to prevent vertical stretching
    # ["S.No", "PMID", "Article Title", "Authors", "Journal", "Abstract"]
    col_widths = {1: 10, 2: 15, 3: 45, 4: 40, 5: 30, 6: 60}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_col = len(fixed_headers) + 1
    
    leaf_col_map = {}

    def write_tree_headers(tree_dict, row, start_col, color_idx=0):
        if not tree_dict:
            return 1 
            
        total_span = 0
        for i, (node_name, children) in enumerate(tree_dict.items()):
            current_color_idx = (color_idx + i) % len(palette) if row == 1 else color_idx
            fill_color = palette[current_color_idx]
            
            child_span = write_tree_headers(children, row + 1, start_col + total_span, current_color_idx)
            
            cell = ws.cell(row=row, column=start_col + total_span, value=node_name)
            cell.font = dynamic_header_font
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            cell.alignment = center_align
            
            if child_span > 1:
                ws.merge_cells(
                    start_row=row, 
                    start_column=start_col + total_span, 
                    end_row=row, 
                    end_column=start_col + total_span + child_span - 1
                )
            else:
                if not children:
                    leaf_col_map[node_name] = start_col + total_span
                    if row < max_depth:
                        ws.merge_cells(
                            start_row=row,
                            start_column=start_col + total_span,
                            end_row=max_depth,
                            end_column=start_col + total_span
                        )
                
            total_span += child_span
            
        return total_span

    write_tree_headers(draft_tree, row=1, start_col=current_col)

    max_col = ws.max_column
    for col_idx in range(len(fixed_headers) + 1, max_col + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 35

    # Fill Rows with PubMed Data
    row_idx = max_depth + 1
    
    order = input_ids if input_ids else pubmed_data.keys()
    
    for pmid in order:
        if pmid not in pubmed_data:
            continue
        p_data = pubmed_data[pmid]
        ws.cell(row=row_idx, column=1, value=row_idx - max_depth).alignment = cell_align
        ws.cell(row=row_idx, column=2, value=pmid).alignment = cell_align
        ws.cell(row=row_idx, column=3, value=p_data.get("title", "N/A")).alignment = cell_align
        ws.cell(row=row_idx, column=4, value=", ".join(p_data.get("authors", []))[:200]).alignment = cell_align
        ws.cell(row=row_idx, column=5, value=p_data.get("journal", "N/A")).alignment = cell_align
        
        abstract_text = p_data.get("abstract", "N/A")
        if len(abstract_text) > 150:
            abstract_text = abstract_text[:150] + "..."
        ws.cell(row=row_idx, column=6, value=abstract_text).alignment = cell_align

        for leaf_name, col_idx in leaf_col_map.items():
            node = next((n for n in final_taxonomy if n["name"] == leaf_name), None)
            if node and pmid in node.get("supporting_pmids", []):
                pmid_specific_contexts = node.get("contexts_by_pmid", {}).get(pmid, [])
                if pmid_specific_contexts:
                    contexts = "\n\n".join(pmid_specific_contexts)
                    ws.cell(row=row_idx, column=col_idx, value=contexts).alignment = cell_align
                else:
                    ws.cell(row=row_idx, column=col_idx, value="✓").alignment = cell_align
            else:
                ws.cell(row=row_idx, column=col_idx, value="-").alignment = center_align

        row_idx += 1

    wb.save(output_path)
