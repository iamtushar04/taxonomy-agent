import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from typing import List, Dict

def get_max_depth(tree_dict, current=0):
    if not tree_dict:
        return current
    return max(get_max_depth(v, current + 1) for v in tree_dict.values() if isinstance(v, dict))

def generate_dynamic_matrix_excel(patents_data: dict, final_taxonomy: List[Dict], draft_tree: dict, output_path: str, input_ids: List[str] = None):
    """
    Generates a dynamically structured Matrix Excel file where:
    - Rows are Patents (in exact user-provided order if input_ids is provided)
    - Columns use merged, multi-level hierarchical headers from draft_tree (e.g. Technology -> Heating -> Heater Design)
    - Cells contain the exact quotes extracted from the patent for the specific leaf node.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patent Taxonomy Matrix"

    fixed_headers = ["S.No", "Patent Number", "Patent Title", "Assignee", "Summary"]
    
    fixed_header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Dark Blue
    fixed_header_font = Font(bold=True, color="FFFFFF")
    
    dynamic_header_font = Font(bold=True, color="000000")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    palette = ["92D050", "00B0F0", "FFC000", "B1A0C7", "FFFF00", "B7DEE8", "D9E1F2", "FCE4D6"]

    max_depth = get_max_depth(draft_tree)
    if max_depth == 0:
        max_depth = 1

    # Write fixed headers
    for col_idx, header in enumerate(fixed_headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = fixed_header_font
        cell.fill = fixed_header_fill
        cell.alignment = center_align
        if max_depth > 1:
            ws.merge_cells(start_row=1, start_column=col_idx, end_row=max_depth, end_column=col_idx)
            
    # Set specific widths for Patent fixed columns to prevent vertical stretching
    # ["S.No", "Patent Number", "Patent Title", "Assignee", "Summary"]
    col_widths = {1: 10, 2: 20, 3: 45, 4: 40, 5: 60}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    current_col = len(fixed_headers) + 1
    
    # Map leaf node names to column indices so we know where to write data
    leaf_col_map = {}

    def write_tree_headers(tree_dict, row, start_col, color_idx=0):
        if not tree_dict:
            return 1 # Leaf node takes exactly 1 column
            
        total_span = 0
        for i, (node_name, children) in enumerate(tree_dict.items()):
            current_color_idx = (color_idx + i) % len(palette) if row == 1 else color_idx
            fill_color = palette[current_color_idx]
            
            child_span = write_tree_headers(children, row + 1, start_col + total_span, current_color_idx)
            
            cell = ws.cell(row=row, column=start_col + total_span, value=node_name)
            cell.font = dynamic_header_font
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            cell.alignment = center_align
            
            if child_span > 1:
                ws.merge_cells(
                    start_row=row, 
                    start_column=start_col + total_span, 
                    end_row=row, 
                    end_column=start_col + total_span + child_span - 1
                )
            else:
                # If child_span == 1 and no children dict, this is a leaf node!
                if not children:
                    leaf_col_map[node_name] = start_col + total_span
                    # If this leaf is reached before max_depth, merge it down to max_depth
                    if row < max_depth:
                        ws.merge_cells(
                            start_row=row,
                            start_column=start_col + total_span,
                            end_row=max_depth,
                            end_column=start_col + total_span
                        )
                
            total_span += child_span
            
        return total_span

    write_tree_headers(draft_tree, row=1, start_col=current_col)

    # Adjust widths for dynamic columns
    max_col = ws.max_column
    for col_idx in range(len(fixed_headers) + 1, max_col + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 35

    # Fill Rows with Patent Data
    row_idx = max_depth + 1
    
    order = input_ids if input_ids else patents_data.keys()
    
    for patent_id in order:
        if patent_id not in patents_data:
            continue
        p_data = patents_data[patent_id]
        # Fixed Data
        ws.cell(row=row_idx, column=1, value=row_idx - max_depth).alignment = cell_align
        ws.cell(row=row_idx, column=2, value=patent_id).alignment = cell_align
        ws.cell(row=row_idx, column=3, value=p_data.get("title", "N/A")).alignment = cell_align
        ws.cell(row=row_idx, column=4, value=p_data.get("assignee", "N/A")).alignment = cell_align
        
        summary_text = p_data.get("summary", "N/A")
        if len(summary_text) > 150:
            summary_text = summary_text[:150] + "..."
        ws.cell(row=row_idx, column=5, value=summary_text).alignment = cell_align

        # Dynamic Data
        for leaf_name, col_idx in leaf_col_map.items():
            node = next((n for n in final_taxonomy if n["name"] == leaf_name), None)
            if node and patent_id in node.get("supporting_patent_ids", []):
                patent_specific_contexts = node.get("contexts_by_patent", {}).get(patent_id, [])
                if patent_specific_contexts:
                    contexts = "\n\n".join(patent_specific_contexts)
                    ws.cell(row=row_idx, column=col_idx, value=contexts).alignment = cell_align
                else:
                    ws.cell(row=row_idx, column=col_idx, value="✓").alignment = cell_align
            else:
                ws.cell(row=row_idx, column=col_idx, value="-").alignment = center_align

        row_idx += 1

    wb.save(output_path)
