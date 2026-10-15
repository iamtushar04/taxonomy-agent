import sys
import os
from typing import List
# pyrefly: ignore [missing-import]
from fastapi import HTTPException

# Add the parent folder to the Python path so it can find taxonomy_builder
parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
if parent_dir not in sys.path:
    sys.path.append(parent_dir)

# pyrefly: ignore [missing-import]
import models
# pyrefly: ignore [missing-import]
from database import SessionLocal
# pyrefly: ignore [missing-import]
from services import db_service
from taxonomy_builder.utils.tracer import init_trace

def run_taxonomy_pipeline(run_id: str, patent_ids: List[str], session_upload_id: str = None, mental_model_context: str = None):
    try:
        db = SessionLocal()
        db_service.update_run_status_running(db, run_id)
            
        # VERY IMPORTANT FOR 1,000 PATENTS: Close the DB connection BEFORE the 5 hour run!
        # Otherwise PostgreSQL will kill the idle connection and the save will fail!
        db.close()
            
        # Clean up IDs: Users might paste "12345 67890" without commas
        cleaned_ids = []
        for raw_id in patent_ids:
            cleaned_ids.extend([part.strip() for part in raw_id.replace(',', ' ').split() if part.strip()])
        patent_ids = cleaned_ids

        init_trace(run_id, patent_ids)
        
        is_pubmed = all(pid.isdigit() for pid in patent_ids)
        
        # Check if products were uploaded in this session
        products_data = {}
        if session_upload_id:
            try:
                # pyrefly: ignore [missing-import]
                from routes.products import _UPLOAD_CACHE
                cached_list = _UPLOAD_CACHE.get(session_upload_id, [])
                for prod in cached_list:
                    # prod is ParsedProduct model
                    p_dict = prod.dict() if hasattr(prod, 'dict') else dict(prod)
                    products_data[p_dict.get("product_id")] = p_dict
            except Exception as pe:
                print(f"Could not load products for session {session_upload_id}: {pe}")

        if is_pubmed:
            from taxonomy_builder.pubmed_graph import build_pubmed_graph
            graph = build_pubmed_graph()
            initial_state = {
                "run_id": run_id,
                "pmids": patent_ids,
                "all_concepts": [],
                "errors": []
            }
        else:
            from taxonomy_builder.graph import build_taxonomy_graph
            graph = build_taxonomy_graph()
            initial_state = {
                "run_id": run_id,
                "patent_ids": patent_ids,
                "products_data": products_data,
                "all_concepts": [],
                "errors": []
            }
            if mental_model_context:
                initial_state["mental_model_context"] = mental_model_context
            
        # INCREASED RECURSION LIMIT FOR 1,000 PATENTS
        config = {"configurable": {"thread_id": run_id}, "recursion_limit": 2000}
        final_state = graph.invoke(initial_state, config=config)
        
        # Open a completely fresh connection to save the results
        db = SessionLocal()
        
        final_taxonomy = final_state.get("final_taxonomy", [])
        raw_source = final_state.get("pubmed_data", final_state.get("patents_data", {}))
        
        # Structure source_data so both patents and products are persisted
        source_data = dict(raw_source)
        if products_data:
            source_data["_products"] = products_data
        
        # Generate meaningful name
        meaningful_name = None
        if final_taxonomy:
            roots = [n for n in final_taxonomy if not n.get("parent_node_id")]
            if len(roots) > 1:
                meaningful_name = " & ".join([r.get("name", "") for r in roots[:2]]) + " Taxonomy"
            elif len(roots) == 1:
                root = roots[0]
                root_name = root.get("name", "")
                children = [n for n in final_taxonomy if n.get("parent_node_id") == root.get("node_id")]
                generic_roots = {"technology", "technologies", "innovation", "innovations", "patent", "patents", "invention", "inventions", "concept", "concepts"}
                
                if not children:
                    meaningful_name = f"{root_name} Taxonomy"
                else:
                    child_names = [c.get("name", "") for c in children[:2]]
                    if root_name.lower() in generic_roots:
                        meaningful_name = " & ".join(child_names) + " Taxonomy"
                    else:
                        meaningful_name = " & ".join(child_names) + f" ({root_name})"
        
        db_service.update_run_status_completed(db, run_id, final_taxonomy, source_data, is_pubmed, meaningful_name)

        # Automatically post taxonomy completion announcement to ChatSession
        try:
            import datetime
            chat_sess = db.query(models.ChatSession).filter(models.ChatSession.run_id == run_id).first()
            if not chat_sess:
                import uuid
                run_obj = db.query(models.Run).filter(models.Run.id == run_id).first()
                chat_sess = models.ChatSession(
                    id=str(uuid.uuid4()),
                    user_id=run_obj.user_id if run_obj and run_obj.user_id else 1,
                    run_id=run_id,
                    messages=[]
                )
                db.add(chat_sess)
                db.commit()

            if chat_sess:
                key_branches = []
                for n in (final_taxonomy or []):
                    if n.get("level") in (1, 2) and n.get("name") not in ["Technology", "Technologies", "Patent", "Patents"]:
                        p_count = len(n.get("supporting_patent_ids") or [])
                        p_str = f" ({p_count} patents)" if p_count > 0 else ""
                        key_branches.append(f"• **{n.get('name')}**{p_str}")
                    if len(key_branches) >= 4:
                        break

                prod_context = ""
                if chat_sess.mental_model and isinstance(chat_sess.mental_model, dict):
                    pname = chat_sess.mental_model.get("product_name")
                    if pname:
                        prod_context = f" aligned with your focus on **{pname}**"

                branches_md = "\n".join(key_branches) if key_branches else "• Comprehensive technological hierarchy"

                completion_msg = (
                    f"🌿 **Taxonomy Generated Successfully!**\n\n"
                    f"I have analyzed your **{len(patent_ids)} item(s)** and structured them into a **{len(final_taxonomy)}-category taxonomy**{prod_context}.\n\n"
                    f"**Key Hierarchy Branches:**\n{branches_md}\n\n"
                    f"The full interactive graph is ready on the right! You can ask me:\n"
                    f"• *'Which patents focus on specific product capabilities?'*\n"
                    f"• *'Why are patents grouped under these categories?'*\n"
                    f"• *'Map [Concept X] under [Concept Y]'* to customize branches."
                )

                msgs = list(chat_sess.messages or [])
                msgs.append({
                    "role": "assistant",
                    "content": completion_msg,
                    "intent": "taxonomy_completed",
                    "metadata": {
                        "type": "taxonomy_completion_summary",
                        "node_count": len(final_taxonomy or []),
                        "patent_count": len(patent_ids),
                        "run_name": meaningful_name
                    },
                    "timestamp": datetime.datetime.utcnow().isoformat()
                })
                chat_sess.messages = msgs
                db.commit()
        except Exception as sess_err:
            print(f"[taxonomy_service] Could not append completion message for run {run_id}: {sess_err}")
    except Exception as e:
        db_service.update_run_status_failed(db, run_id, str(e))
    finally:
        try:
            # pyrefly: ignore [missing-import]
            from services.run_stage_service import clear_run_stage
            clear_run_stage(run_id)
        except Exception:
            pass
        db.close()


def process_excel_generation(run_id: str, payload, run_record):
    is_pubmed = run_record.is_pubmed
    source_data = run_record.source_data or {}
    
    parent_map = {e.target: e.source for e in payload.edges}
    flat_nodes = []
    
    for n in payload.nodes:
        pmids = set(n.data.supporting_pmids) if n.data.supporting_pmids else set()
        if not pmids and n.data.contexts_by_pmid:
            pmids = set(n.data.contexts_by_pmid.keys())
            
        patents = set(n.data.supporting_patent_ids) if n.data.supporting_patent_ids else set()
        if not patents and n.data.contexts_by_patent:
            patents = set(n.data.contexts_by_patent.keys())

        products = set(n.data.supporting_product_ids) if n.data.supporting_product_ids else set()
        if not products and getattr(n.data, "contexts_by_product", None):
            products = set(n.data.contexts_by_product.keys())
        
        flat_nodes.append({
            "node_id": n.id,
            "parent_node_id": parent_map.get(n.id),
            "name": n.data.label,
            "level": 0, 
            "description": "",
            "supporting_pmids": list(pmids),
            "supporting_patent_ids": list(patents),
            "supporting_product_ids": list(products),
            "contexts_by_patent": n.data.contexts_by_patent or {},
            "contexts_by_pmid": n.data.contexts_by_pmid or {},
            "contexts_by_product": getattr(n.data, "contexts_by_product", None) or {}
        })
        
    children_map = {}
    for n in flat_nodes:
        if n["parent_node_id"]:
            children_map.setdefault(n["parent_node_id"], []).append(n["node_id"])
            
    node_map = {n["node_id"]: n for n in flat_nodes}
    
    def collect_descendants(node_id, visited):
        if node_id in visited: return set(), set(), set()
        visited.add(node_id)
        res_pmid = set(node_map[node_id].get("supporting_pmids", []))
        res_pat = set(node_map[node_id].get("supporting_patent_ids", []))
        res_prod = set(node_map[node_id].get("supporting_product_ids", []))
        for child in children_map.get(node_id, []):
            cp, ct, cpr = collect_descendants(child, visited)
            res_pmid.update(cp)
            res_pat.update(ct)
            res_prod.update(cpr)
        return res_pmid, res_pat, res_prod
        
    for n in flat_nodes:
        p1, p2, p3 = collect_descendants(n["node_id"], set())
        n["supporting_pmids"] = sorted(list(p1))
        n["supporting_patent_ids"] = sorted(list(p2))
        n["supporting_product_ids"] = sorted(list(p3))
        
    roots = [n for n in flat_nodes if not n["parent_node_id"]]
    def _build(node_id):
        children = [n for n in flat_nodes if n["parent_node_id"] == node_id]
        if not children: return {}
        return {c["name"]: _build(c["node_id"]) for c in children}
        
    enriched_tree = {root["name"]: _build(root["node_id"]) for root in roots}
    
    matrix_path = os.path.join(parent_dir, "data", f"taxonomy_matrix_{run_id}.xlsx")
    
    # Clean up input_ids in case they were pasted with spaces instead of commas
    cleaned_input_ids = []
    if run_record.input_ids:
        for raw_id in run_record.input_ids:
            cleaned_input_ids.extend([part.strip() for part in raw_id.replace(',', ' ').split() if part.strip()])
    else:
        cleaned_input_ids = None

    # Load products data
    products_data = source_data.get("_products")
    if not products_data:
        try:
            # pyrefly: ignore [missing-import]
            from database import SessionLocal
            # pyrefly: ignore [missing-import]
            import models
            s_db = SessionLocal()
            chat_sess = s_db.query(models.ChatSession).filter(models.ChatSession.run_id == run_id).first()
            if chat_sess and chat_sess.mental_model:
                mm_prods = chat_sess.mental_model.get("products")
                if mm_prods:
                    products_data = {p.get("product_id", f"prod_{i}"): p for i, p in enumerate(mm_prods)}
            s_db.close()
        except Exception as e:
            print(f"Could not load products from chat session: {e}")

    if is_pubmed:
        from taxonomy_builder.utils.pubmed_excel_matrix_generator import generate_pubmed_matrix_excel
        generate_pubmed_matrix_excel(source_data, flat_nodes, enriched_tree, matrix_path, cleaned_input_ids)
    else:
        from taxonomy_builder.utils.excel_matrix_generator import generate_dynamic_matrix_excel
        generate_dynamic_matrix_excel(source_data, flat_nodes, enriched_tree, matrix_path, cleaned_input_ids, products_data=products_data)
        
    if not os.path.exists(matrix_path):
        raise HTTPException(status_code=404, detail="Excel file generation failed.")
        
    return matrix_path


def build_excel_preview_data(matrix_path: str) -> dict:
    """
    Parses the generated taxonomy matrix Excel file into structured JSON
    including merged cells, colors, fonts, and dimensions so the frontend can
    render an exact pixel-matched preview before downloading.
    Supports multi-sheet workbooks.
    """
    import openpyxl
    
    if not os.path.exists(matrix_path):
        raise HTTPException(status_code=404, detail="Excel matrix file not found.")

    wb = openpyxl.load_workbook(matrix_path, data_only=True)
    sheet_names = wb.sheetnames

    def parse_sheet(ws):
        merged_ranges = []
        covered_cells = set()
        for m in ws.merged_cells.ranges:
            merged_ranges.append({
                'min_row': m.min_row,
                'max_row': m.max_row,
                'min_col': m.min_col,
                'max_col': m.max_col,
                'rowspan': m.max_row - m.min_row + 1,
                'colspan': m.max_col - m.min_col + 1
            })
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    if not (r == m.min_row and c == m.min_col):
                        covered_cells.add((r, c))

        merge_by_top_left = {(m['min_row'], m['min_col']): m for m in merged_ranges}

        rows = []
        max_r = ws.max_row or 1
        max_c = ws.max_column or 1

        col_widths = {}
        for c in range(1, max_c + 1):
            col_letter = openpyxl.utils.get_column_letter(c)
            dim = ws.column_dimensions.get(col_letter)
            col_widths[c] = dim.width if dim and dim.width else 18

        for r in range(1, max_r + 1):
            row_cells = []
            for c in range(1, max_c + 1):
                if (r, c) in covered_cells:
                    continue

                cell = ws.cell(r, c)
                val = cell.value

                fill_color = None
                if cell.fill and cell.fill.start_color and cell.fill.start_color.rgb:
                    rgb_str = str(cell.fill.start_color.rgb)
                    if len(rgb_str) == 8 and rgb_str != '00000000':
                        fill_color = '#' + rgb_str[2:]
                    elif len(rgb_str) == 6:
                        fill_color = '#' + rgb_str

                font_color = None
                if cell.font and cell.font.color and cell.font.color.rgb:
                    c_str = str(cell.font.color.rgb)
                    if len(c_str) == 8 and c_str != '00000000':
                        font_color = '#' + c_str[2:]
                    elif len(c_str) == 6:
                        font_color = '#' + c_str

                bold = cell.font.bold if cell.font else False
                align_h = cell.alignment.horizontal if cell.alignment else 'left'
                align_v = cell.alignment.vertical if cell.alignment else 'middle'

                cell_data = {
                    'row': r,
                    'col': c,
                    'value': val if val is not None else '',
                    'bg': fill_color,
                    'color': font_color,
                    'bold': bold,
                    'align_h': align_h,
                    'align_v': align_v,
                }

                if (r, c) in merge_by_top_left:
                    cell_data['rowspan'] = merge_by_top_left[(r, c)]['rowspan']
                    cell_data['colspan'] = merge_by_top_left[(r, c)]['colspan']

                row_cells.append(cell_data)
            rows.append(row_cells)

        return {
            "title": ws.title,
            "max_row": max_r,
            "max_col": max_c,
            "col_widths": col_widths,
            "rows": rows
        }

    sheets_data = {}
    for name in sheet_names:
        sheets_data[name] = parse_sheet(wb[name])

    active_name = sheet_names[0] if sheet_names else ""
    first_data = sheets_data.get(active_name, {"title": "", "max_row": 0, "max_col": 0, "col_widths": {}, "rows": []})

    return {
        "sheet_names": sheet_names,
        "active_sheet": active_name,
        "title": first_data["title"],
        "max_row": first_data["max_row"],
        "max_col": first_data["max_col"],
        "col_widths": first_data["col_widths"],
        "rows": first_data["rows"],
        "sheets": sheets_data
    }
